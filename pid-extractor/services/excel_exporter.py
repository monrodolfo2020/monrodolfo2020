import logging
from collections import Counter
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from models.instrument import ExtractionResult, Instrument

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Colour constants
# ---------------------------------------------------------------------------
_NAVY = "1F3864"
_WHITE = "FFFFFF"
_LIGHT_GRAY = "F2F2F2"
_GREEN_DARK = "375623"
_RED_DARK = "C00000"
_GRAY_DARK = "7F7F7F"

_FILL_RED = PatternFill(fill_type="solid", fgColor="FFCCCC")
_FILL_YELLOW = PatternFill(fill_type="solid", fgColor="FFF2CC")
_FILL_GREEN_LIGHT = PatternFill(fill_type="solid", fgColor="CCFFCC")
_FILL_WHITE = PatternFill(fill_type="solid", fgColor=_WHITE)
_FILL_GRAY = PatternFill(fill_type="solid", fgColor=_LIGHT_GRAY)
_FILL_NAVY_HDR = PatternFill(fill_type="solid", fgColor=_NAVY)

_FONT_HEADER = Font(bold=True, color=_WHITE, size=11)
_ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
_ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _set_col_widths(ws, widths: List[float]) -> None:
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def _header_cell(ws, row: int, col: int, value: str) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = _FONT_HEADER
    cell.fill = _FILL_NAVY_HDR
    cell.alignment = _ALIGN_CENTER


# ---------------------------------------------------------------------------
# Sheet 1 — Instruments
# ---------------------------------------------------------------------------
_INSTR_HEADERS = [
    "Tag", "Measured Variable", "Function", "Loop Number",
    "Line/Pipe", "P&ID Number", "System Title",
    "Confidence", "Notes", "Page",
]
_INSTR_WIDTHS = [15, 22, 18, 14, 28, 18, 38, 12, 42, 8]


def _build_instruments_sheet(ws, instruments: List[Instrument]) -> None:
    ws.title = "Instruments"
    ws.sheet_properties.tabColor = _NAVY

    # Header row
    for col, header in enumerate(_INSTR_HEADERS, start=1):
        _header_cell(ws, 1, col, header)
    ws.row_dimensions[1].height = 22

    # Data rows
    for row_idx, inst in enumerate(instruments, start=2):
        row_data = [
            inst.tag,
            inst.measured_variable,
            inst.function_letters,
            inst.loop_number,
            inst.line_pipe,
            inst.pid_number,
            inst.system_title,
            inst.confidence.upper() if inst.confidence else "",
            inst.notes,
            inst.page_number,
        ]

        is_even = (row_idx % 2 == 0)
        base_fill = _FILL_WHITE if is_even else _FILL_GRAY
        conf = (inst.confidence or "high").lower()

        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = base_fill
            cell.alignment = _ALIGN_CENTER if col_idx in (8, 10) else _ALIGN_LEFT

        # Confidence-based colouring on Tag (col 1) and Confidence (col 8)
        if conf == "low":
            ws.cell(row=row_idx, column=1).fill = _FILL_RED
            ws.cell(row=row_idx, column=8).fill = _FILL_RED
        elif conf == "medium":
            ws.cell(row=row_idx, column=1).fill = _FILL_YELLOW
            ws.cell(row=row_idx, column=8).fill = _FILL_YELLOW
        else:  # high
            ws.cell(row=row_idx, column=8).fill = _FILL_GREEN_LIGHT

    # Auto-filter + freeze pane
    last_row = max(len(instruments) + 1, 1)
    last_col = get_column_letter(len(_INSTR_HEADERS))
    ws.auto_filter.ref = f"A1:{last_col}{last_row}"
    ws.freeze_panes = "A2"

    _set_col_widths(ws, _INSTR_WIDTHS)


# ---------------------------------------------------------------------------
# Sheet 2 — Summary
# ---------------------------------------------------------------------------
def _bold_cell(ws, row: int, col: int, value, size: int = 11) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(bold=True, size=size)
    return cell


def _section_header(ws, row: int, label: str) -> None:
    cell = ws.cell(row=row, column=1, value=label)
    cell.font = Font(bold=True, size=11, color=_NAVY)


def _build_summary_sheet(ws, result: ExtractionResult) -> None:
    ws.title = "Summary"
    ws.sheet_properties.tabColor = _GREEN_DARK

    # Title
    ws.merge_cells("A1:B1")
    title_cell = ws.cell(row=1, column=1, value="P&ID EXTRACTION SUMMARY")
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = _ALIGN_CENTER

    # Key/value pairs starting at row 3
    kv_pairs = [
        ("Source File:", result.source_pdf),
        ("Extraction Date:", result.created_at.strftime("%Y-%m-%d %H:%M UTC")),
        ("Total Pages Analyzed:", result.total_pages),
        ("Total Instruments Found:", len(result.instruments)),
        ("Errors:", len(result.errors)),
        ("Warnings:", len(result.warnings)),
    ]
    current_row = 3
    for label, value in kv_pairs:
        ws.cell(row=current_row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=current_row, column=2, value=value)
        current_row += 1

    current_row += 1  # blank spacer

    # --- By measured variable ---
    _section_header(ws, current_row, "INSTRUMENTS BY MEASURED VARIABLE")
    current_row += 1
    var_counts = Counter(
        inst.measured_variable or "Unknown" for inst in result.instruments
    )
    for var, count in sorted(var_counts.items(), key=lambda x: -x[1]):
        ws.cell(row=current_row, column=1, value=var)
        ws.cell(row=current_row, column=2, value=count)
        current_row += 1

    current_row += 1

    # --- By P&ID number ---
    _section_header(ws, current_row, "INSTRUMENTS BY P&ID NUMBER")
    current_row += 1
    pid_counts = Counter(
        inst.pid_number or "(unknown)" for inst in result.instruments
    )
    for pid, count in sorted(pid_counts.items(), key=lambda x: -x[1]):
        ws.cell(row=current_row, column=1, value=pid)
        ws.cell(row=current_row, column=2, value=count)
        current_row += 1

    current_row += 1

    # --- By confidence level ---
    _section_header(ws, current_row, "INSTRUMENTS BY CONFIDENCE LEVEL")
    current_row += 1
    conf_counts: Dict[str, int] = {"high": 0, "medium": 0, "low": 0}
    for inst in result.instruments:
        key = (inst.confidence or "high").lower()
        conf_counts[key] = conf_counts.get(key, 0) + 1
    for level in ("high", "medium", "low"):
        ws.cell(row=current_row, column=1, value=level.upper())
        ws.cell(row=current_row, column=2, value=conf_counts[level])
        current_row += 1

    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 20


# ---------------------------------------------------------------------------
# Sheet 3 — Extraction Log
# ---------------------------------------------------------------------------
_LOG_HEADERS = ["Page", "Status", "Instruments Found", "Notes/Errors"]
_LOG_WIDTHS = [8, 12, 20, 60]


def _build_log_sheet(ws, result: ExtractionResult) -> None:
    ws.title = "Extraction Log"

    has_errors = bool(result.errors)
    ws.sheet_properties.tabColor = _RED_DARK if has_errors else _GRAY_DARK

    # Header row
    for col, header in enumerate(_LOG_HEADERS, start=1):
        _header_cell(ws, 1, col, header)
    ws.row_dimensions[1].height = 22

    # Per-page log
    # Build helper maps
    instr_by_page: Dict[int, List[Instrument]] = {}
    for inst in result.instruments:
        instr_by_page.setdefault(inst.page_number, []).append(inst)

    # Errors are stored as strings like "Page N: message"
    error_pages: set = set()
    for err in result.errors:
        # Try to detect page number from error message prefix "Page N:"
        if err.lower().startswith("page "):
            parts = err.split(":", 1)
            try:
                pnum = int(parts[0].split()[1])
                error_pages.add(pnum)
            except (IndexError, ValueError):
                pass

    for page in range(1, result.total_pages + 1):
        row = page + 1
        page_instruments = instr_by_page.get(page, [])
        low_conf = [i.tag for i in page_instruments if i.confidence == "low"]

        if page in error_pages:
            status = "ERROR"
            fill = _FILL_RED
        elif low_conf:
            status = "WARNING"
            fill = _FILL_YELLOW
        else:
            status = "OK"
            fill = _FILL_GREEN_LIGHT

        notes = ""
        if low_conf:
            notes = "Low confidence: " + ", ".join(low_conf)

        row_data = [page, status, len(page_instruments), notes]
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = _ALIGN_CENTER if col_idx <= 3 else _ALIGN_LEFT

    # Error detail section
    if result.errors:
        error_start = result.total_pages + 3
        ws.cell(row=error_start, column=1, value="ERRORS").font = Font(bold=True, color="C00000")
        for idx, err in enumerate(result.errors, start=1):
            ws.cell(row=error_start + idx, column=1, value=idx)
            ws.cell(row=error_start + idx, column=2, value=err)

    _set_col_widths(ws, _LOG_WIDTHS)


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------
def export_to_excel(result: ExtractionResult, output_path: str) -> str:
    """Build a 3-sheet Excel workbook and save it to *output_path*.

    Returns the output_path string after saving.
    """
    wb = Workbook()

    # openpyxl creates a default sheet; rename/reuse for Instruments
    ws_instruments = wb.active
    _build_instruments_sheet(ws_instruments, result.instruments)

    ws_summary = wb.create_sheet("Summary")
    _build_summary_sheet(ws_summary, result)

    ws_log = wb.create_sheet("Extraction Log")
    _build_log_sheet(ws_log, result)

    wb.properties.creator = "PID Extractor"
    wb.properties.title = "P&ID Instrument Extraction"

    wb.save(output_path)
    logger.info("Excel workbook saved to '%s'", output_path)
    return output_path
